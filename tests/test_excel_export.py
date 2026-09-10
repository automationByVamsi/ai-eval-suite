"""Tests for PO Excel export."""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from src.models.evaluation_result import (
    CaseEvaluationResult,
    DeterministicCheckResult,
    E2ECaseResult,
)
from src.models.metric_result import MetricResult
from src.reporting.excel_export import export_evaluation_report, export_filename


def _stage_result(
    *,
    test_case_id: str = "TC_001",
    passed_det: bool = True,
    passed_judge: bool = True,
) -> CaseEvaluationResult:
    return CaseEvaluationResult(
        eval_name="sanity",
        test_case_id=test_case_id,
        agent_name="knowledge_agent",
        question="How do I support someone gambling?",
        answer="Encourage professional help.",
        expected_output="Seek professional help.",
        latency_ms=1200.0,
        run_id="20260823_120000",
        deterministic_results=[
            DeterministicCheckResult(
                name="answer_non_empty",
                passed=passed_det,
                reason="" if passed_det else "empty answer",
            )
        ],
        metric_results=[
            MetricResult(
                name="intent_preservation",
                score=0.92 if passed_judge else 0.4,
                threshold=0.7,
                passed=passed_judge,
                reason="" if passed_judge else "intent drift",
            )
        ],
        result_fields={
            "rewritten_query": "gambling support",
            "anchor_page_id": "9001",
            "business_area": "/Support_Hub/",
            "decision": "FINISH",
        },
    )


def _load_sheets(data: bytes) -> dict[str, list[list]]:
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    out: dict[str, list[list]] = {}
    for name in wb.sheetnames:
        ws = wb[name]
        out[name] = [list(row) for row in ws.iter_rows(values_only=True)]
    wb.close()
    return out


def test_export_filename():
    assert export_filename(run_id="20260823_120000", view_mode="stage").endswith(".xlsx")


def test_export_stage_workbook_sheets_and_summary():
    stage = _stage_result()
    data = export_evaluation_report(
        stage_results=[stage],
        view_mode="stage",
        scope_label="Single run · 20260823_120000",
        run_id="20260823_120000",
    )
    sheets = _load_sheets(data)

    assert list(sheets) == [
        "Executive Summary",
        "Test Results",
        "Checks Detail",
        "Failed Cases",
        "Case Metadata",
    ]
    summary = {row[0]: row[1] for row in sheets["Executive Summary"] if row[0]}
    assert summary["Test cases"] == "1"
    assert summary["Overall pass rate"] == "100%"

    test_rows = sheets["Test Results"]
    assert test_rows[1][0] == "TC_001"
    assert test_rows[1][4] == "PASS"

    checks = sheets["Checks Detail"]
    assert len(checks) == 3  # header + det + judge
    assert checks[1][4] == "deterministic"
    assert checks[2][4] == "judge"

    failed = sheets["Failed Cases"]
    assert failed[1][0] == "No failed cases in this export scope."


def test_export_includes_failed_case():
    stage = _stage_result(passed_judge=False)
    data = export_evaluation_report(stage_results=[stage], view_mode="stage")
    sheets = _load_sheets(data)

    assert sheets["Test Results"][1][4] == "FAIL"
    assert sheets["Failed Cases"][1][0] == "TC_001"
    assert "intent drift" in str(sheets["Failed Cases"][1][7])


def test_export_e2e_view():
    stage = _stage_result()
    e2e = E2ECaseResult(
        test_case_id="TC_001",
        agent_name="knowledge_agent",
        question=stage.question,
        latency_ms=stage.latency_ms,
        stages=[stage],
        run_id="20260823_120000",
    )
    data = export_evaluation_report(e2e_results=[e2e], view_mode="e2e")
    sheets = _load_sheets(data)

    assert sheets["Test Results"][1][2] == "e2e (1 stages)"
    assert sheets["Test Results"][1][4] == "PASS"


def test_export_requires_valid_view_mode():
    with pytest.raises(ValueError, match="view_mode"):
        export_evaluation_report(view_mode="invalid")
